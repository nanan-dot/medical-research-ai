"""Office document parsers with explicit legacy DOC conversion boundaries."""

import shutil
import subprocess
from collections.abc import Iterable
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Protocol, cast

from docx import Document as WordDocument
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide

from app.modules.document.parsers.base import (
    DocumentParserError,
    LegacyDocumentConverterUnavailableError,
    validate_extracted_size,
    validate_file_size,
)
from app.modules.document.parsers.schemas import (
    ParsedDocument,
    ParsedPage,
    ParsedSection,
)


def _non_empty_lines(values: list[str]) -> list[str]:
    return [value.strip() for value in values if value and value.strip()]


class TextCell(Protocol):
    @property
    def text(self) -> str: ...


class TextRow(Protocol):
    @property
    def cells(self) -> Iterable[TextCell]: ...


class TextTable(Protocol):
    @property
    def rows(self) -> Iterable[TextRow]: ...


def _table_lines(tables: Iterable[TextTable]) -> list[str]:
    lines: list[str] = []
    for table in tables:
        for row in table.rows:
            values = _non_empty_lines([cell.text for cell in row.cells])
            if values:
                lines.append(" | ".join(values))
    return lines


class DocxParser:
    """Extract text and heading-style sections from a modern Word document."""

    def parse(self, path: Path) -> ParsedDocument:
        validate_file_size(path)
        try:
            document = WordDocument(str(path))
        except Exception as exc:
            raise DocumentParserError("Word document extraction failed") from exc

        paragraphs = _non_empty_lines(
            [paragraph.text for paragraph in document.paragraphs]
        )
        table_lines = _table_lines(document.tables)
        text = "\n".join([*paragraphs, *table_lines])
        validate_extracted_size(text)
        title = (document.core_properties.title or "").strip()
        sections = [
            ParsedSection(heading=paragraph.text.strip(), level=1, text="")
            for paragraph in document.paragraphs
            if (
                paragraph.text.strip()
                and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")
            )
        ]
        return ParsedDocument(
            source_path=str(path),
            title=(title or next(iter(paragraphs), path.stem))[:500],
            text=text,
            sections=sections,
        )


class PptxParser:
    """Extract visible slide text while retaining one-based slide locations."""

    def parse(self, path: Path) -> ParsedDocument:
        validate_file_size(path)
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise DocumentParserError("PowerPoint extraction failed") from exc

        pages = [
            ParsedPage(page_number=index, text=self._slide_text(slide))
            for index, slide in enumerate(presentation.slides, start=1)
        ]
        text = "\n\n".join(page.text for page in pages if page.text)
        validate_extracted_size(text)
        title = (presentation.core_properties.title or "").strip()
        first_line = next(
            (line for page in pages for line in page.text.splitlines() if line),
            path.stem,
        )
        return ParsedDocument(
            source_path=str(path),
            title=(title or first_line)[:500],
            text=text,
            pages=pages,
        )

    @staticmethod
    def _slide_text(slide: Slide) -> str:
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_shape = cast(Shape, shape)
                lines.extend(
                    _non_empty_lines(
                        [
                            paragraph.text
                            for paragraph in text_shape.text_frame.paragraphs
                        ]
                    )
                )
            if shape.has_table:
                table_shape = cast(GraphicFrame, shape)
                lines.extend(_table_lines([table_shape.table]))
        return "\n".join(lines)


class LegacyDocParser:
    """Convert legacy DOC with an installed local converter, never cloud upload."""

    _CONVERTERS = ("antiword", "catdoc")

    def parse(self, path: Path) -> ParsedDocument:
        validate_file_size(path)
        text = self._extract_text(path)
        validate_extracted_size(text)
        return ParsedDocument(
            source_path=str(path),
            title=next((line for line in text.splitlines() if line), path.stem)[:500],
            text=text,
        )

    def _extract_text(self, path: Path) -> str:
        for converter in self._CONVERTERS:
            executable = shutil.which(converter)
            if executable:
                return self._run_converter([executable, str(path)])
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            return self._convert_with_libreoffice(soffice, path)
        raise LegacyDocumentConverterUnavailableError(
            "Legacy .doc parsing requires LibreOffice, antiword, or catdoc"
        )

    @staticmethod
    def _run_converter(command: list[str]) -> str:
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                check=True,
                encoding="utf-8",
                errors="replace",
                timeout=60,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            raise DocumentParserError("Legacy Word document conversion failed") from exc
        return result.stdout.strip()

    @staticmethod
    def _convert_with_libreoffice(executable: str, path: Path) -> str:
        with TemporaryDirectory() as output_directory:
            command = [
                executable,
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                output_directory,
                str(path),
            ]
            LegacyDocParser._run_converter(command)
            converted = Path(output_directory) / f"{path.stem}.txt"
            try:
                return converted.read_text(encoding="utf-8", errors="replace").strip()
            except OSError as exc:
                raise DocumentParserError(
                    "Legacy Word document conversion failed"
                ) from exc
