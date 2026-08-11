"""将 DOCX 转换为供安全插值渲染的只读结构化内容。"""

from __future__ import annotations

from typing import Any

from docx import Document as WordDocument

from app.modules.document_preview.schema import (
    DocumentPreviewBlock,
    DocumentPreviewBlockKind,
    DocumentPreviewTable,
)

_MAX_BLOCKS = 2000
_MAX_TABLES = 100
_MAX_ROWS_PER_TABLE = 200
_MAX_CELLS_PER_ROW = 50
_MAX_TEXT_LENGTH = 10000


def convert_docx_to_preview(path: str) -> tuple[list[DocumentPreviewBlock], list[DocumentPreviewTable]]:
    """只输出文本数据而非原始 HTML，避免 Office 内容跨越信任边界。"""

    document = WordDocument(path)
    blocks = _convert_paragraphs(document)
    tables = _convert_tables(document)
    return blocks, tables


def _convert_paragraphs(document: Any) -> list[DocumentPreviewBlock]:
    blocks: list[DocumentPreviewBlock] = []
    for paragraph in document.paragraphs:
        text = _compact_text(paragraph.text)
        if not text:
            continue
        blocks.append(
            DocumentPreviewBlock(
                kind=_paragraph_kind(paragraph.style.name),
                level=_heading_level(paragraph.style.name),
                text=text,
            )
        )
        if len(blocks) >= _MAX_BLOCKS:
            break
    return blocks


def _convert_tables(document: Any) -> list[DocumentPreviewTable]:
    tables: list[DocumentPreviewTable] = []
    for table in document.tables[:_MAX_TABLES]:
        rows = [
            [_compact_text(cell.text) for cell in row.cells[:_MAX_CELLS_PER_ROW]]
            for row in table.rows[:_MAX_ROWS_PER_TABLE]
        ]
        if rows:
            tables.append(DocumentPreviewTable(rows=rows))
    return tables


def _paragraph_kind(style_name: str) -> DocumentPreviewBlockKind:
    lowered_style = style_name.casefold()
    if lowered_style.startswith("heading"):
        return DocumentPreviewBlockKind.HEADING
    if "list" in lowered_style:
        return DocumentPreviewBlockKind.LIST_ITEM
    return DocumentPreviewBlockKind.PARAGRAPH


def _heading_level(style_name: str) -> int | None:
    if not style_name.casefold().startswith("heading"):
        return None
    suffix = style_name.removeprefix("Heading").strip()
    return int(suffix) if suffix.isdigit() and 1 <= int(suffix) <= 9 else 1


def _compact_text(value: str) -> str:
    return " ".join(value.split())[:_MAX_TEXT_LENGTH]
