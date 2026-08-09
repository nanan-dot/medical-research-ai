from pathlib import Path

from docx import Document as WordDocument
from pptx import Presentation

from app.modules.document.parsers.base import LegacyDocumentConverterUnavailableError
from app.modules.document.parsers.factory import create_parser
from app.modules.document.parsers.office_parser import LegacyDocParser


def test_docx_parser_extracts_title_text_and_heading(tmp_path: Path):
    path = tmp_path / "study.docx"
    document = WordDocument()
    document.core_properties.title = "Cardiology Study"
    document.add_heading("Methods", level=1)
    document.add_paragraph("A prospective cohort was analysed.")
    document.save(path)

    parsed = create_parser(path).parse(path)

    assert parsed.title == "Cardiology Study"
    assert "prospective cohort" in parsed.text
    assert [(section.heading, section.level) for section in parsed.sections] == [
        ("Methods", 1)
    ]


def test_pptx_parser_preserves_slide_locations(tmp_path: Path):
    path = tmp_path / "report.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Weekly Literature Review"
    slide.placeholders[1].text = "EGFR evidence and study limitations"
    presentation.save(path)

    parsed = create_parser(path).parse(path)

    assert parsed.title == "Weekly Literature Review"
    assert len(parsed.pages) == 1
    assert parsed.pages[0].page_number == 1
    assert "study limitations" in parsed.pages[0].text


def test_legacy_doc_requires_local_converter(tmp_path: Path, monkeypatch):
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy binary document")
    monkeypatch.setattr("shutil.which", lambda _: None)

    parser = LegacyDocParser()
    try:
        parser.parse(path)
    except LegacyDocumentConverterUnavailableError as error:
        assert "LibreOffice" in str(error)
    else:
        raise AssertionError("A missing legacy converter must not be silently ignored")
