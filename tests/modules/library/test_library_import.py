"""Acceptance coverage for multi-format library import safety."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfWriter


def _pdf_bytes() -> BytesIO:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    content = BytesIO()
    writer.write(content)
    content.seek(0)
    return content


def _docx_bytes() -> BytesIO:
    document = WordDocument()
    document.add_paragraph("DOCX fixture")
    content = BytesIO()
    document.save(content)
    content.seek(0)
    return content


def _pptx_bytes() -> BytesIO:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    content = BytesIO()
    presentation.save(content)
    content.seek(0)
    return content


def test_multi_format_import_creates_document_tasks(library_client) -> None:
    """AC-LIB-07: supported formats are independently accepted and queued."""
    client, _, _ = library_client
    response = client.post(
        "/api/v1/library/imports",
        files=[
            ("files", ("paper.pdf", _pdf_bytes(), "application/pdf")),
            ("files", ("notes.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("slides.pptx", _pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
            ("files", ("notes.md", BytesIO(b"# Markdown"), "text/markdown")),
            ("files", ("plain.txt", BytesIO(b"Plain text"), "text/plain")),
        ],
    )

    assert response.status_code == 201
    items = response.json()["items"]
    assert [item["status"] for item in items] == ["accepted"] * 5
    assert all(item["task_id"] is not None for item in items)


def test_import_rejects_unsafe_inputs_and_deduplicates(library_client, monkeypatch) -> None:
    """AC-LIB-08: unsafe inputs leave no managed file and duplicate hash is deterministic."""
    client, _, upload_root = library_client
    first = client.post(
        "/api/v1/library/imports",
        files=[("files", ("same.md", BytesIO(b"same content"), "text/markdown"))],
    )
    assert first.status_code == 201
    duplicate = client.post(
        "/api/v1/library/imports",
        files=[("files", ("same.md", BytesIO(b"same content"), "text/markdown"))],
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["items"][0]["status"] == "duplicate"

    unsafe = client.post(
        "/api/v1/library/imports",
        files=[
            ("files", ("../escape.md", BytesIO(b"text"), "text/markdown")),
            ("files", ("fake.pdf", BytesIO(b"not a PDF"), "application/pdf")),
        ],
    )
    assert unsafe.status_code == 201
    assert [item["status"] for item in unsafe.json()["items"]] == ["rejected", "rejected"]
    assert not list(upload_root.rglob("escape.md"))

    monkeypatch.setattr("app.core.config.settings.MAX_LIBRARY_ARCHIVE_UNCOMPRESSED_BYTES", 32)
    archive = BytesIO()
    with ZipFile(archive, "w", ZIP_DEFLATED) as zip_file:
        zip_file.writestr("word/document.xml", "x" * 128)
        zip_file.writestr("[Content_Types].xml", "types")
    archive.seek(0)
    bomb = client.post(
        "/api/v1/library/imports",
        files=[("files", ("too-large.docx", archive, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert bomb.status_code == 201
    assert bomb.json()["items"][0]["error_code"] == "archive_too_large"
