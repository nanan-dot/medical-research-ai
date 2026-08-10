import asyncio
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core import models  # noqa: F401
from app.core.database import Base, get_session
from app.main import app
from app.modules.document.model import Document
from app.modules.knowledge_source.model import KnowledgeSource


@pytest.fixture
def client(tmp_path: Path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{(tmp_path / 'preview.db').as_posix()}")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    root = tmp_path / "source"
    root.mkdir()
    pdf_path = root / "paper.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf_path.open("wb") as stream:
        writer.write(stream)
    docx_path = root / "notes.docx"
    word_document = WordDocument()
    word_document.add_heading("研究摘要", level=1)
    word_document.add_paragraph("<script>not executable</script>")
    word_document.save(docx_path)
    legacy_doc_path = root / "legacy.doc"
    legacy_doc_path.write_bytes(b"legacy word content")
    outside_path = tmp_path / "outside.pdf"
    outside_path.write_bytes(pdf_path.read_bytes())

    async def prepare() -> dict[str, int]:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with session_factory() as session:
            source = KnowledgeSource(
                name="preview source",
                source_type="local_folder",
                root_path=str(root),
                normalized_root_path=str(root).casefold(),
                enabled=True,
                sync_status="idle",
            )
            session.add(source)
            await session.flush()
            identifiers: dict[str, int] = {}
            for label, file_path in {
                "pdf": pdf_path,
                "docx": docx_path,
                "doc": legacy_doc_path,
                "outside": outside_path,
            }.items():
                stat = file_path.stat()
                document = Document(
                    knowledge_source_id=source.id,
                    file_path="../outside.pdf" if label == "outside" else file_path.name,
                    normalized_file_path=f"{label}/{file_path.name}",
                    file_hash="a" * 64,
                    file_size=stat.st_size,
                    modified_time=datetime.fromtimestamp(stat.st_mtime, UTC),
                    modified_time_ns=stat.st_mtime_ns,
                    scan_state="pending",
                    parse_status="pending",
                    index_status="pending",
                )
                session.add(document)
                await session.flush()
                identifiers[label] = document.id
            await session.commit()
            return identifiers

    async def override_session():
        async with session_factory() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    identifiers = asyncio.run(prepare())
    app.dependency_overrides[get_session] = override_session
    with TestClient(app) as test_client:
        yield test_client, identifiers
    app.dependency_overrides.clear()
    asyncio.run(engine.dispose())


def test_pdf_preview_descriptor_and_inline_range_stream(client):
    test_client, document_ids = client

    preview = test_client.get(f"/api/v1/documents/{document_ids['pdf']}/preview")
    assert preview.status_code == 200
    assert preview.json() == {
        "document_id": document_ids["pdf"],
        "kind": "pdf",
        "content_url": f"/api/v1/documents/{document_ids['pdf']}/original",
        "blocks": [],
        "tables": [],
        "message": None,
    }

    streamed = test_client.get(
        f"/api/v1/documents/{document_ids['pdf']}/original",
        headers={"Range": "bytes=0-4"},
    )
    assert streamed.status_code == 206
    assert streamed.content == b"%PDF-"
    assert streamed.headers["content-type"].startswith("application/pdf")
    assert streamed.headers["content-disposition"].startswith("inline")
    assert streamed.headers["x-content-type-options"] == "nosniff"
    assert streamed.headers["accept-ranges"] == "bytes"


def test_pdf_range_suffix_and_out_of_bounds(client):
    test_client, document_ids = client
    pdf_id = document_ids["pdf"]

    # 后缀范围 "bytes=-5"：返回文件最后 5 字节（Content-Range 含 end 与总大小）
    suffix = test_client.get(
        f"/api/v1/documents/{pdf_id}/original",
        headers={"Range": "bytes=-5"},
    )
    assert suffix.status_code == 206
    assert len(suffix.content) == 5
    assert "/" in suffix.headers["content-range"]

    # 越界范围：start 超过文件大小 → 416 + Content-Range: bytes */size
    out_of_bounds = test_client.get(
        f"/api/v1/documents/{pdf_id}/original",
        headers={"Range": "bytes=999999-1000000"},
    )
    assert out_of_bounds.status_code == 416
    assert out_of_bounds.headers["content-range"].startswith("bytes */")

    # 无 Range 头：完整文件 200
    full = test_client.get(f"/api/v1/documents/{pdf_id}/original")
    assert full.status_code == 200
    assert full.content.startswith(b"%PDF-")


def test_docx_preview_returns_structured_content_without_raw_html(client):
    test_client, document_ids = client

    response = test_client.get(f"/api/v1/documents/{document_ids['docx']}/preview")

    assert response.status_code == 200
    payload = response.json()
    assert payload["kind"] == "docx"
    assert payload["content_url"] is None
    assert payload["blocks"][0] == {"kind": "heading", "text": "研究摘要", "level": 1}
    assert payload["blocks"][1]["text"] == "<script>not executable</script>"


def test_unsupported_or_out_of_root_documents_do_not_leak_files(client):
    test_client, document_ids = client

    legacy = test_client.get(f"/api/v1/documents/{document_ids['doc']}/preview")
    assert legacy.status_code == 200
    assert legacy.json()["kind"] == "unavailable"

    escaped = test_client.get(f"/api/v1/documents/{document_ids['outside']}/preview")
    assert escaped.status_code == 409
    assert escaped.json()["error"]["code"] == "document_preview_unavailable"

    missing = test_client.get("/api/v1/documents/99999/preview")
    assert missing.status_code == 404
